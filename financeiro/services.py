import re
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from django.db import transaction
from .models import Lancamento

def _campo(bloco, nome):
    match = re.search(rf"<{nome}>([^<\r\n]+)", bloco, re.I)
    return match.group(1).strip() if match else ""

@transaction.atomic
def importar_ofx(arquivo, conta, categoria_entrada, categoria_despesa, usuario, centro_custo=None):
    texto = arquivo.read().decode("utf-8", errors="ignore")
    blocos = re.findall(r"<STMTTRN>(.*?)(?:</STMTTRN>|(?=<STMTTRN>|</BANKTRANLIST>))", texto, re.I | re.S)
    criados = ignorados = 0
    for bloco in blocos:
        valor_texto = _campo(bloco, "TRNAMT").replace(",", ".")
        data_texto = _campo(bloco, "DTPOSTED")[:8]
        if not valor_texto or len(data_texto) != 8: continue
        valor_assinado = Decimal(valor_texto)
        fitid_original = _campo(bloco, "FITID") or f"{data_texto}:{valor_texto}:{_campo(bloco, 'NAME')}"
        fitid = f"conta-{conta.pk}:{fitid_original}"
        if Lancamento.objects.filter(identificador_bancario=fitid).exists():
            ignorados += 1; continue
        entrada = valor_assinado >= 0
        Lancamento.objects.create(descricao=(_campo(bloco, "MEMO") or _campo(bloco, "NAME") or "Movimento bancário")[:160], tipo="E" if entrada else "D", valor=abs(valor_assinado), data=datetime.strptime(data_texto, "%Y%m%d").date(), categoria=categoria_entrada if entrada else categoria_despesa, conta=conta, centro_custo=centro_custo, documento=_campo(bloco, "CHECKNUM")[:60], identificador_bancario=fitid, conciliado=True, criado_por=usuario)
        criados += 1
    return criados, ignorados

def relatorio_xlsx(lancamentos):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = "Prestação de contas"
    cabecalho = ["Data", "Tipo", "Descrição", "Categoria", "Centro de custo", "Conta", "Documento", "Valor"]
    ws.append(cabecalho)
    for cell in ws[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="087D3E")
    for item in lancamentos:
        ws.append([item.data, item.get_tipo_display(), item.descricao, item.categoria.nome, item.centro_custo.nome if item.centro_custo else "", item.conta.nome, item.documento, float(item.valor if item.tipo == "E" else -item.valor)])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for col, largura in {"A":12,"B":12,"C":35,"D":24,"E":24,"F":20,"G":18,"H":16}.items(): ws.column_dimensions[col].width = largura
    output = BytesIO(); wb.save(output); return output.getvalue()

def relatorio_pdf(lancamentos, titulo):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    output = BytesIO(); doc = SimpleDocTemplate(output, pagesize=landscape(A4), title=titulo)
    styles = getSampleStyleSheet(); dados = [["Data", "Tipo", "Descrição", "Categoria", "Conta", "Valor"]]
    total = Decimal("0")
    for item in lancamentos:
        sinal = item.valor if item.tipo == "E" else -item.valor; total += sinal
        dados.append([item.data.strftime("%d/%m/%Y"), item.get_tipo_display(), item.descricao[:45], item.categoria.nome, item.conta.nome, f"R$ {sinal:,.2f}"])
    tabela = Table(dados, repeatRows=1, colWidths=[65,60,230,130,120,90])
    tabela.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#087D3E")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),.4,colors.HexColor("#dfe7e2")),("ALIGN",(-1,1),(-1,-1),"RIGHT"),("FONTSIZE",(0,0),(-1,-1),8),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#f4f7f5")])]))
    doc.build([Paragraph(titulo, styles["Title"]), Spacer(1,12), tabela, Spacer(1,12), Paragraph(f"Saldo do período: R$ {total:,.2f}", styles["Heading2"])]); return output.getvalue()
